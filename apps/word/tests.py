import io

from django.test import SimpleTestCase
from docx import Document

from .views import html_to_docx_bytes, sanitize_html


class WordExportTests(SimpleTestCase):
    def test_sanitize_preserves_editor_formatting_and_removes_unsafe_css(self):
        source = (
            '<p style="text-align:center;line-height:1.5">'
            '<span style="font-family:Arial;font-size:14pt;font-weight:bold;'
            'color:#112233;background-image:url(https://bad.example)">Текст</span>'
            '<script>alert(1)</script></p>'
        )
        clean = sanitize_html(source)
        self.assertIn('text-align: center', clean)
        self.assertIn('line-height: 1.5', clean)
        self.assertIn('font-family: arial', clean)
        self.assertIn('font-size: 14pt', clean)
        self.assertIn('font-weight: bold', clean)
        self.assertIn('color: #112233', clean)
        self.assertNotIn('background-image', clean)
        self.assertNotIn('script', clean)

    def test_export_preserves_run_styles_and_page_size(self):
        body = (
            '<p style="text-align:center">'
            '<span style="font-size:14pt;font-weight:bold;color:#112233">Готово</span>'
            '</p>'
        )
        data = html_to_docx_bytes(sanitize_html(body), 'A5')
        doc = Document(io.BytesIO(data))
        run = doc.paragraphs[0].runs[0]
        self.assertEqual(run.text, 'Готово')
        self.assertTrue(run.bold)
        self.assertEqual(run.font.size.pt, 14)
        self.assertEqual(str(run.font.color.rgb), '112233')
        self.assertAlmostEqual(doc.sections[0].page_width.mm, 148, delta=0.1)
        self.assertAlmostEqual(doc.sections[0].page_height.mm, 210, delta=0.1)

    def test_export_embeds_local_images(self):
        pixel = (
            'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk'
            '+A8AAQUBAScY42YAAAAASUVORK5CYII='
        )
        body = f'<p><img src="data:image/png;base64,{pixel}" style="width:30px"></p>'
        doc = Document(io.BytesIO(html_to_docx_bytes(sanitize_html(body))))
        self.assertEqual(len(doc.inline_shapes), 1)

    def test_export_decodes_table_text(self):
        doc = Document(io.BytesIO(html_to_docx_bytes(
            '<table><tr><td>А &amp; Б</td></tr></table>'
        )))
        self.assertEqual(doc.tables[0].cell(0, 0).text, 'А & Б')
